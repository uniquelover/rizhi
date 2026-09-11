#!/usr/bin/env python3
from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "out" / "vehicle-settings-analysis"
PURE_UI_CSV = OUT_DIR / "车辆设置_纯UI用例总表.csv"
MASTER_CSV = OUT_DIR / "vehicle_settings_control_mapping.csv"


HEADERS = [
    "模块",
    "控件名",
    "控件类型",
    "一级页面",
    "二级页面",
    "页面名",
    "是否已抓图",
    "图片路径",
    "图片来源组",
    "关联用例数",
    "关联用例",
    "来源Sheet",
    "是否依赖CAN",
    "备注",
]


COORDINATE_CONTROLS = {
    "自动落锁",
    "走近解锁",
    "远离闭锁",
    "锁车关闭遮阳帘",
    "锁车升窗",
    "驻车解锁",
}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def split_csv_field(value: str) -> list[str]:
    if not value:
        return []
    return [item.strip() for item in value.split(",") if item.strip()]


def build_master_index() -> dict[str, dict[str, str]]:
    rows = read_csv(MASTER_CSV)
    return {row["control_name"]: row for row in rows}


def collect_required_controls() -> list[dict[str, str]]:
    pure_rows = read_csv(PURE_UI_CSV)
    master_index = build_master_index()

    case_map: dict[str, set[str]] = defaultdict(set)
    sheet_map: dict[str, set[str]] = defaultdict(set)
    can_map: dict[str, set[str]] = defaultdict(set)

    for row in pure_rows:
        if row.get("是否纯UI") not in {"是", "ÊÇ"}:
            continue
        case_name = row.get("用例序号") or row.get("用例名") or row.get("功能序号") or ""
        sheet_name = row.get("来源Sheet", "")
        can_flag = row.get("是否依赖CAN", "")
        for control_name in split_csv_field(row.get("匹配控件", "")):
            case_map[control_name].add(case_name)
            if sheet_name:
                sheet_map[control_name].add(sheet_name)
            if can_flag:
                can_map[control_name].add(can_flag)

    out_rows: list[dict[str, str]] = []
    for control_name in sorted(case_map):
        meta = master_index.get(control_name, {})
        out_rows.append(
            {
                "模块": "车辆设置",
                "控件名": control_name,
                "控件类型": "坐标型" if control_name in COORDINATE_CONTROLS else "图片型",
                "一级页面": meta.get("level1", ""),
                "二级页面": meta.get("level2", ""),
                "页面名": meta.get("page_name", ""),
                "是否已抓图": "是" if (meta.get("has_image") == "yes" or control_name in COORDINATE_CONTROLS) else "否",
                "图片路径": meta.get("image_path", ""),
                "图片来源组": meta.get("image_group", ""),
                "关联用例数": str(len(case_map[control_name])),
                "关联用例": ", ".join(sorted(case_map[control_name])),
                "来源Sheet": ", ".join(sorted(sheet_map.get(control_name, set()))),
                "是否依赖CAN": "是" if "是" in can_map.get(control_name, set()) or "ÊÇ" in can_map.get(control_name, set()) else "否",
                "备注": meta.get("notes", ""),
            }
        )
    return out_rows


def style_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.append(HEADERS)
    widths = [10, 22, 12, 12, 14, 18, 10, 56, 18, 12, 42, 18, 12, 30]
    for idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]


def write_xlsx(rows: list[dict[str, str]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "required_controls"
    style_sheet(ws)
    row_no = 2
    for row in rows:
        for col_no, field in enumerate(HEADERS, 1):
            ws.cell(row=row_no, column=col_no, value=row.get(field, ""))
            ws.cell(row=row_no, column=col_no).alignment = Alignment(vertical="center", wrap_text=True)
        row_no += 1
    out_path = OUT_DIR / "车辆设置_可做用例所需控件总表.xlsx"
    wb.save(out_path)
    return out_path


def write_csv(rows: list[dict[str, str]]) -> Path:
    out_path = OUT_DIR / "车辆设置_可做用例所需控件总表.csv"
    with out_path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(rows)
    return out_path


def main() -> None:
    rows = collect_required_controls()
    xlsx_path = write_xlsx(rows)
    csv_path = write_csv(rows)
    print(f"xlsx={xlsx_path.relative_to(ROOT).as_posix()}")
    print(f"csv={csv_path.relative_to(ROOT).as_posix()}")
    print(f"rows={len(rows)}")


if __name__ == "__main__":
    main()
