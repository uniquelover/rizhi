#!/usr/bin/env python3
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "out" / "vehicle-settings-analysis"
SOURCE_CSV = OUT_DIR / "vehicle_settings_control_mapping.csv"


FIELDS = [
    "page_key",
    "level1",
    "level2",
    "page_name",
    "control_name",
    "usage_type",
    "used_by_cases",
    "has_image",
    "image_group",
    "image_path",
    "notes",
]


def load_rows() -> list[dict[str, str]]:
    with SOURCE_CSV.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    return [
        row
        for row in rows
        if row.get("scope_status") == "in_scope" and row.get("has_image") == "no"
        and row.get("usage_type") != "ptouch"
    ]


def write_csv(rows: list[dict[str, str]]) -> Path:
    out_path = OUT_DIR / "vehicle_settings_missing_images.csv"
    with out_path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows({field: row.get(field, "") for field in FIELDS} for row in rows)
    return out_path


def style_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.append(FIELDS)
    widths = [18, 10, 14, 18, 24, 16, 42, 10, 14, 56, 32]
    for idx, header in enumerate(FIELDS, 1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]


def write_xlsx(rows: list[dict[str, str]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "missing_images"
    style_sheet(ws)
    row_no = 2
    for row in rows:
        for col_no, field in enumerate(FIELDS, 1):
            ws.cell(row=row_no, column=col_no, value=row.get(field, ""))
            ws.cell(row=row_no, column=col_no).alignment = Alignment(vertical="center", wrap_text=True)
        row_no += 1
    out_path = OUT_DIR / "vehicle_settings_missing_images.xlsx"
    wb.save(out_path)
    return out_path


def write_md(rows: list[dict[str, str]]) -> Path:
    out_path = OUT_DIR / "vehicle_settings_missing_images.md"
    lines = [
        "# 车辆设置缺图清单",
        "",
        f"- 缺图总数：{len(rows)}",
        "",
        "| 一级模块 | 二级页面 | 控件名 | 用途 | 关联用例 | 备注 |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for row in rows:
        lines.append(
            f"| {row.get('level1', '')} | {row.get('level2', '')} | {row.get('control_name', '')} | "
            f"{row.get('usage_type', '')} | {row.get('used_by_cases', '')} | {row.get('notes', '')} |"
        )
    out_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return out_path


def main() -> None:
    rows = load_rows()
    csv_path = write_csv(rows)
    xlsx_path = write_xlsx(rows)
    md_path = write_md(rows)
    print(f"xlsx={xlsx_path.relative_to(ROOT).as_posix()}")
    print(f"csv={csv_path.relative_to(ROOT).as_posix()}")
    print(f"md={md_path.relative_to(ROOT).as_posix()}")
    print(f"rows={len(rows)}")


if __name__ == "__main__":
    main()
