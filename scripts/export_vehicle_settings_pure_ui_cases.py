#!/usr/bin/env python3
from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "out" / "vehicle-settings-analysis"
MASTER_CSV = OUT_DIR / "vehicle_settings_control_mapping.csv"
WORKBOOK_PATH = Path.home() / "Downloads" / "T1V全功能测试用例 (1).xlsx"
TARGET_SHEETS = ["车辆设置", "能量管理"]


HEADERS = [
    "模块",
    "来源Sheet",
    "功能序号",
    "用例序号",
    "一级功能",
    "二级功能",
    "三级功能",
    "功能点",
    "前提条件",
    "操作步骤",
    "期望结果",
    "是否纯UI",
    "是否依赖CAN",
    "匹配控件",
    "已抓图控件数",
    "缺图控件数",
    "缺图控件",
    "备注",
]


SOURCE_HEADERS = [
    "功能序号",
    "参考需求版本",
    "用例序号",
    "ECU控制器/件",
    "一级功能",
    "二级功能",
    "三级功能",
    "功能点",
    "用例属性",
    "用例方法",
    "前提条件",
    "操作步骤",
    "期望结果",
]


CAN_PATTERNS = [
    r"\bICC_",
    r"\bVCU_",
    r"\bBMS_",
    r"\bFLZCU_",
    r"\bHCU_",
    r"\b0x[0-9A-Fa-f]+",
    r"持续发送",
    r"反馈信号",
    r"CAN",
    r"档位",
    r"插枪",
    r"充电枪",
    r"高配双电机",
    r"配置字",
    r"信号",
]


PURE_UI_NEGATIVE_PATTERNS = [
    r"恢复出厂设置",
    r"扫码",
    r"登录",
    r"账号",
    r"语音",
    r"实车",
    r"蓝牙钥匙",
    r"UWB",
]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def load_cases() -> list[dict[str, str]]:
    wb = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True, keep_links=False)
    rows: list[dict[str, str]] = []
    for sheet_name in TARGET_SHEETS:
        ws = wb[sheet_name]
        for row_idx, row_values in enumerate(
            ws.iter_rows(min_row=2, max_col=len(SOURCE_HEADERS), values_only=True),
            start=2,
        ):
            values = [normalize_text(value) for value in row_values]
            if not any(values):
                continue
            record = dict(zip(SOURCE_HEADERS, values))
            if not record["用例序号"]:
                continue
            record["来源Sheet"] = sheet_name
            rows.append(record)
    return rows


def build_control_catalog() -> list[dict[str, object]]:
    rows = read_csv(MASTER_CSV)
    catalog: list[dict[str, object]] = []
    for row in rows:
        aliases = [item.strip() for item in row.get("control_name", "").split(",") if item.strip()]
        for alias in [item.strip() for item in row.get("image_path", "").split(",") if item.strip()]:
            stem = Path(alias).stem
            if stem and stem not in aliases:
                aliases.append(stem)
        for alias in [item.strip() for item in row.get("page_name", "").split(",") if item.strip()]:
            if alias and alias not in aliases:
                aliases.append(alias)
        catalog.append(
            {
                "control_name": row.get("control_name", ""),
                "aliases": [item for item in aliases if item],
                "has_image": row.get("has_image", "") == "yes",
                "page_name": row.get("page_name", ""),
            }
        )
    return catalog


def match_controls(text_blob: str, catalog: list[dict[str, object]]) -> list[dict[str, object]]:
    matched = []
    for item in catalog:
        aliases: list[str] = item["aliases"]  # type: ignore[assignment]
        if any(alias and alias in text_blob for alias in aliases):
            matched.append(item)
    unique = {}
    for item in matched:
        unique[item["control_name"]] = item
    return list(unique.values())


def infer_can(text_blob: str) -> bool:
    return any(re.search(pattern, text_blob, flags=re.IGNORECASE) for pattern in CAN_PATTERNS)


def infer_pure_ui(text_blob: str, can_dependent: bool) -> str:
    if any(re.search(pattern, text_blob) for pattern in PURE_UI_NEGATIVE_PATTERNS):
        return "否"
    if can_dependent:
        return "是"
    return "是"


def build_rows() -> list[dict[str, str]]:
    cases = load_cases()
    catalog = build_control_catalog()
    out_rows: list[dict[str, str]] = []

    for case in cases:
        text_blob = "\n".join(
            [
                case.get("前提条件", ""),
                case.get("操作步骤", ""),
                case.get("期望结果", ""),
                case.get("二级功能", ""),
                case.get("三级功能", ""),
                case.get("功能点", ""),
            ]
        )
        controls = match_controls(text_blob, catalog)
        control_names = [item["control_name"] for item in controls]
        missing_controls = [item["control_name"] for item in controls if not item["has_image"]]
        can_dependent = infer_can(text_blob)
        pure_ui = infer_pure_ui(text_blob, can_dependent)

        notes = []
        if not controls:
            notes.append("未匹配到既有控件映射")
        if can_dependent:
            notes.append("原始用例文本含信号/CAN动作，但按当前口径仍计入可做范围")

        out_rows.append(
            {
                "模块": "车辆设置",
                "来源Sheet": case["来源Sheet"],
                "功能序号": case["功能序号"],
                "用例序号": case["用例序号"],
                "一级功能": case["一级功能"],
                "二级功能": case["二级功能"],
                "三级功能": case["三级功能"],
                "功能点": case["功能点"],
                "前提条件": case["前提条件"],
                "操作步骤": case["操作步骤"],
                "期望结果": case["期望结果"],
                "是否纯UI": pure_ui,
                "是否依赖CAN": "是" if can_dependent else "否",
                "匹配控件": ", ".join(control_names),
                "已抓图控件数": str(sum(1 for item in controls if item["has_image"])),
                "缺图控件数": str(len(missing_controls)),
                "缺图控件": ", ".join(missing_controls),
                "备注": "；".join(notes),
            }
        )
    return out_rows


def style_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.append(HEADERS)
    widths = [10, 12, 16, 14, 12, 16, 18, 26, 34, 34, 42, 10, 12, 32, 12, 12, 28, 30]
    for idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]


def write_xlsx(rows: list[dict[str, str]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "纯UI用例总表"
    style_sheet(ws)
    row_no = 2
    for row in rows:
        for col_no, field in enumerate(HEADERS, 1):
            ws.cell(row=row_no, column=col_no, value=row.get(field, ""))
            ws.cell(row=row_no, column=col_no).alignment = Alignment(vertical="center", wrap_text=True)
        row_no += 1
    candidates = [
        OUT_DIR / "车辆设置_纯UI用例总表.xlsx",
        OUT_DIR / "车辆设置_纯UI用例总表_v2.xlsx",
        OUT_DIR / "车辆设置_纯UI用例总表_v3.xlsx",
    ]
    last_error = None
    for out_path in candidates:
        try:
            wb.save(out_path)
            return out_path
        except PermissionError as exc:
            last_error = exc
            continue
    raise last_error


def write_csv(rows: list[dict[str, str]]) -> Path:
    candidates = [
        OUT_DIR / "车辆设置_纯UI用例总表.csv",
        OUT_DIR / "车辆设置_纯UI用例总表_v2.csv",
        OUT_DIR / "车辆设置_纯UI用例总表_v3.csv",
    ]
    last_error = None
    for out_path in candidates:
        try:
            with out_path.open("w", encoding="utf-8-sig", newline="") as fh:
                writer = csv.DictWriter(fh, fieldnames=HEADERS)
                writer.writeheader()
                writer.writerows(rows)
            return out_path
        except PermissionError as exc:
            last_error = exc
            continue
    raise last_error


def main() -> None:
    rows = build_rows()
    xlsx_path = write_xlsx(rows)
    csv_path = write_csv(rows)
    print(f"xlsx={xlsx_path.relative_to(ROOT).as_posix()}")
    print(f"csv={csv_path.relative_to(ROOT).as_posix()}")
    print(f"rows={len(rows)}")


if __name__ == "__main__":
    main()
