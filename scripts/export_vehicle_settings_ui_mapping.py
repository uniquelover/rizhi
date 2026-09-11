#!/usr/bin/env python3
from __future__ import annotations

import csv
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from config.vehicle_settings_icon_mapping_config import (
    CURATED_DIR_HINTS,
    LABEL_SEARCH_DIRS,
    LABEL_ALIASES,
    MANUAL_LABEL_PATHS,
    PAGE_TREE,
    ROOT,
    SCRIPT_INCLUDE_KEYWORDS,
)

OUT_DIR = ROOT / "out" / "vehicle-settings-analysis"


def ensure_out_dir() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)


def decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "cp936"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def iter_vehicle_script_files() -> list[Path]:
    script_files = []
    for path in (ROOT / "press_android").rglob("*.txt"):
        rel = path.relative_to(ROOT).as_posix()
        if any(keyword in rel for keyword in SCRIPT_INCLUDE_KEYWORDS):
            script_files.append(path)
    return sorted(script_files)


def normalize_label(label: str) -> str:
    label = label.strip()
    if not label:
        return ""
    if "%" in label:
        label = label.split("%", 1)[0].strip()
    return label.strip()


def split_combo_labels(payload: str) -> list[str]:
    parts = [item for item in payload.strip().split() if item]
    return [normalize_label(item) for item in parts[:2] if normalize_label(item)]


def parse_script(path: Path) -> list[dict[str, str]]:
    text = decode_text(path.read_bytes())
    rows: list[dict[str, str]] = []
    step_no = 0
    for line in text.splitlines():
        line = line.strip()
        if not line or ":" not in line or line.startswith("#"):
            continue
        command, payload = line.split(":", 1)
        command = command.strip()
        payload = payload.strip()
        labels: list[tuple[str, str]] = []
        if command == "touch":
            label = normalize_label(payload)
            if label:
                labels.append(("touch", label))
        elif command == "check":
            label = normalize_label(payload)
            if label:
                labels.append(("check", label))
        elif command == "ptouch":
            label = normalize_label(payload)
            if label:
                labels.append(("ptouch", label))
        elif command == "ctouch":
            combo = split_combo_labels(payload)
            if len(combo) >= 1:
                labels.append(("ctouch_check", combo[0]))
            if len(combo) >= 2:
                labels.append(("ctouch_touch", combo[1]))
        for action_type, label in labels:
            step_no += 1
            rows.append(
                {
                    "script_file": path.relative_to(ROOT).as_posix(),
                    "case_name": path.stem,
                    "step_no": str(step_no),
                    "action_type": action_type,
                    "label": label,
                }
            )
    return rows


def build_script_usage() -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for path in iter_vehicle_script_files():
        rows.extend(parse_script(path))
    return rows


def build_label_index() -> dict[str, list[Path]]:
    index: dict[str, list[Path]] = defaultdict(list)
    for base_dir in LABEL_SEARCH_DIRS:
        if not base_dir.exists():
            continue
        for path in base_dir.rglob("*"):
            if path.is_file() and path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
                index[path.stem].append(path)
    return dict(index)


def path_group(path: Path) -> str:
    rel = path.relative_to(ROOT).as_posix()
    for prefix, title in CURATED_DIR_HINTS.items():
        if rel.startswith(prefix):
            return title
    return "其他"


def page_record_index() -> dict[str, dict[str, str]]:
    return {item["page_key"]: item for item in PAGE_TREE}


def infer_page_key(label: str) -> str:
    for item in PAGE_TREE:
        if label == item["entry_label"] or label in item["anchor_labels"]:
            return item["page_key"]
    if label in {"Home", "Home1", "T1VHome", "T1VHome1"}:
        return "public_home"
    if "车辆设置icon" in label or "Dock车辆设置" in label or label == "T1VDock车辆设置":
        return "public_dock_vehicle_settings"
    if "常用操作" in label or "后视镜" in label:
        return "vehicle_common"
    if "门窗锁" in label or "车门锁" in label or "解锁" in label or "闭锁" in label or "锁车" in label:
        return "vehicle_door_lock"
    if "车门" in label or "儿童锁" in label or "后备箱" in label or "智能钥匙" in label:
        return "vehicle_door_lock"
    if "充放电" in label or "充电" in label or "能量管理" in label or "能量回收" in label:
        return "energy_charge"
    if "车外灯" in label or "位置灯" in label or "近光灯" in label or "后雾灯" in label:
        return "light_exterior"
    if (
        "亮度" in label
        or "显示" in label
        or "浅色" in label
        or "深色" in label
        or "12小时" in label
        or "24小时" in label
        or label in {"屏幕", "自动new"}
    ):
        return "display_root"
    if "声音" in label or "音效" in label or "声浪" in label or "导航音量" in label:
        return "sound_root"
    if "连接" in label or "CarPlay" in label or "HiCar" in label or "无线" in label:
        return "connection_root"
    if "恢复出厂设置" in label or label.startswith("通用"):
        return "general_root"
    if label in {"T1VDock车辆设置", "Home1"}:
        return "public_dock_vehicle_settings" if label == "T1VDock车辆设置" else "public_home"
    return ""


def resolve_best_image(label: str, label_index: dict[str, list[Path]]) -> Path | None:
    manual_rel = MANUAL_LABEL_PATHS.get(label)
    if manual_rel:
        manual_path = ROOT / manual_rel
        if manual_path.exists():
            return manual_path
    candidates = list(label_index.get(label, []))
    for alias in LABEL_ALIASES.get(label, []):
        candidates.extend(label_index.get(alias, []))
    if not candidates:
        return None
    preferred_roots = [str(path).lower() for path in LABEL_SEARCH_DIRS]
    candidates = sorted(
        candidates,
        key=lambda path: (
            next((idx for idx, root in enumerate(preferred_roots) if str(path).lower().startswith(root)), 999),
            len(path.parts),
            str(path),
        ),
    )
    return candidates[0]


def build_control_mapping(script_rows: list[dict[str, str]], label_index: dict[str, list[Path]]) -> list[dict[str, str]]:
    usage_map: dict[str, set[str]] = defaultdict(set)
    script_map: dict[str, set[str]] = defaultdict(set)
    for row in script_rows:
        usage_map[row["label"]].add(row["action_type"])
        script_map[row["label"]].add(row["case_name"])

    labels = set(usage_map.keys())
    for item in PAGE_TREE:
        for label in [item["entry_label"], *item["anchor_labels"]]:
            if label:
                labels.add(label)

    rows: list[dict[str, str]] = []
    page_idx = page_record_index()
    for label in sorted(labels):
        page_key = infer_page_key(label)
        page_meta = page_idx.get(page_key, {})
        image_path = resolve_best_image(label, label_index)
        rows.append(
            {
                "page_key": page_key,
                "level1": page_meta.get("level1", ""),
                "level2": page_meta.get("level2", ""),
                "page_name": page_meta.get("page_name", ""),
                "scope_status": "in_scope" if page_key else "historical_or_unmapped",
                "control_name": label,
                "usage_type": "/".join(sorted(usage_map.get(label, set()))),
                "used_by_cases": ", ".join(sorted(script_map.get(label, set()))),
                "has_image": "yes" if image_path else "no",
                "image_path": image_path.relative_to(ROOT).as_posix() if image_path else "",
                "image_group": path_group(image_path) if image_path else "",
                "notes": page_meta.get("notes", ""),
            }
        )
    return rows


def build_page_tree_rows(control_rows: list[dict[str, str]], script_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    control_count: dict[str, int] = defaultdict(int)
    case_count: dict[str, set[str]] = defaultdict(set)
    for row in control_rows:
        if row["page_key"]:
            control_count[row["page_key"]] += 1
            if row["used_by_cases"]:
                case_count[row["page_key"]].update(item.strip() for item in row["used_by_cases"].split(",") if item.strip())
    page_rows = []
    for item in PAGE_TREE:
        page_rows.append(
            {
                "page_key": item["page_key"],
                "level1": item["level1"],
                "level2": item["level2"],
                "page_name": item["page_name"],
                "entry_label": item["entry_label"],
                "anchor_labels": ", ".join(item["anchor_labels"]),
                "control_count": str(control_count.get(item["page_key"], 0)),
                "script_case_count": str(len(case_count.get(item["page_key"], set()))),
                "notes": item["notes"],
            }
        )
    return page_rows


def build_historical_summary(label_index: dict[str, list[Path]]) -> list[dict[str, str]]:
    grouped: dict[str, list[Path]] = defaultdict(list)
    for paths in label_index.values():
        for path in paths:
            grouped[path_group(path)].append(path)
    rows = []
    for group, paths in sorted(grouped.items()):
        rows.append(
            {
                "source_group": group,
                "image_count": str(len(paths)),
                "sample_paths": ", ".join(path.relative_to(ROOT).as_posix() for path in sorted(paths)[:5]),
            }
        )
    return rows


def write_csv(path: Path, rows: list[dict[str, str]], fieldnames: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def style_sheet(ws, headers: list[str], widths: list[int]) -> None:
    ws.freeze_panes = "A2"
    ws.append(headers)
    for idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]


def write_sheet_rows(ws, rows: list[dict[str, str]], fieldnames: list[str]) -> None:
    row_no = 2
    for row in rows:
        for col_no, field in enumerate(fieldnames, 1):
            ws.cell(row=row_no, column=col_no, value=row.get(field, ""))
            ws.cell(row=row_no, column=col_no).alignment = Alignment(vertical="center", wrap_text=True)
        row_no += 1


def write_workbook(page_rows: list[dict[str, str]], control_rows: list[dict[str, str]], script_rows: list[dict[str, str]], history_rows: list[dict[str, str]]) -> Path:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    ws_pages = wb.create_sheet("page_tree")
    page_fields = ["page_key", "level1", "level2", "page_name", "entry_label", "anchor_labels", "control_count", "script_case_count", "notes"]
    style_sheet(ws_pages, page_fields, [18, 12, 14, 18, 18, 32, 12, 14, 36])
    write_sheet_rows(ws_pages, page_rows, page_fields)

    ws_controls = wb.create_sheet("control_mapping")
    control_fields = ["page_key", "level1", "level2", "page_name", "scope_status", "control_name", "usage_type", "used_by_cases", "has_image", "image_group", "image_path", "notes"]
    style_sheet(ws_controls, control_fields, [18, 10, 14, 18, 18, 22, 18, 28, 10, 18, 56, 32])
    write_sheet_rows(ws_controls, control_rows, control_fields)

    ws_scripts = wb.create_sheet("script_usage")
    script_fields = ["script_file", "case_name", "step_no", "action_type", "label"]
    style_sheet(ws_scripts, script_fields, [52, 28, 10, 18, 24])
    write_sheet_rows(ws_scripts, script_rows, script_fields)

    ws_history = wb.create_sheet("historical_assets")
    history_fields = ["source_group", "image_count", "sample_paths"]
    style_sheet(ws_history, history_fields, [20, 12, 100])
    write_sheet_rows(ws_history, history_rows, history_fields)

    out_path = OUT_DIR / "vehicle_settings_ui_mapping.xlsx"
    wb.save(out_path)
    return out_path


def write_summary_md(page_rows: list[dict[str, str]], control_rows: list[dict[str, str]], script_rows: list[dict[str, str]], history_rows: list[dict[str, str]]) -> Path:
    lines = [
        "# 车辆设置页面树与控件映射",
        "",
        "## 页面树",
        "",
        "| 一级模块 | 二级页面 | 页面名 | 入口锚点 | 已映射控件数 | 脚本覆盖用例数 |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for row in page_rows:
        lines.append(
            f"| {row['level1']} | {row['level2']} | {row['page_name']} | {row['entry_label']} | {row['control_count']} | {row['script_case_count']} |"
        )
    lines.extend(
        [
            "",
            "## 控件映射说明",
            "",
            f"- 控件总数：{len(control_rows)}",
            f"- 脚本动作总数：{len(script_rows)}",
            f"- 已有图控件数：{sum(1 for row in control_rows if row['has_image'] == 'yes')}",
            f"- 缺图控件数：{sum(1 for row in control_rows if row['has_image'] == 'no')}",
            "",
            "## 历史资源概览",
            "",
            "| 资源组 | 图片数 | 样例路径 |",
            "| --- | --- | --- |",
        ]
    )
    for row in history_rows:
        lines.append(f"| {row['source_group']} | {row['image_count']} | {row['sample_paths']} |")
    out_path = OUT_DIR / "vehicle_settings_ui_mapping.md"
    out_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return out_path


def main() -> None:
    ensure_out_dir()
    script_rows = build_script_usage()
    label_index = build_label_index()
    control_rows = build_control_mapping(script_rows, label_index)
    page_rows = build_page_tree_rows(control_rows, script_rows)
    history_rows = build_historical_summary(label_index)

    write_csv(OUT_DIR / "vehicle_settings_page_tree.csv", page_rows, ["page_key", "level1", "level2", "page_name", "entry_label", "anchor_labels", "control_count", "script_case_count", "notes"])
    write_csv(OUT_DIR / "vehicle_settings_control_mapping.csv", control_rows, ["page_key", "level1", "level2", "page_name", "scope_status", "control_name", "usage_type", "used_by_cases", "has_image", "image_group", "image_path", "notes"])
    write_csv(OUT_DIR / "vehicle_settings_script_usage.csv", script_rows, ["script_file", "case_name", "step_no", "action_type", "label"])
    workbook_path = write_workbook(page_rows, control_rows, script_rows, history_rows)
    summary_path = write_summary_md(page_rows, control_rows, script_rows, history_rows)

    print(f"xlsx={workbook_path.relative_to(ROOT).as_posix()}")
    print(f"md={summary_path.relative_to(ROOT).as_posix()}")
    print(f"pages={len(page_rows)}")
    print(f"controls={len(control_rows)}")
    print(f"script_actions={len(script_rows)}")


if __name__ == "__main__":
    main()
