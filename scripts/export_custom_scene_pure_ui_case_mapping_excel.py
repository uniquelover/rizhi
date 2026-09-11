#!/usr/bin/env python3
import csv
import importlib.util
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage


ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "out" / "custom-scene-analysis"
ANALYSIS_CSV = OUT_DIR / "analysis.csv"
DETAIL_CSV = OUT_DIR / "ui_control_mapping_detail.csv"
CONFIG_PATH = ROOT / "config" / "custom_scene_ui_pipeline_config.py"

PURE_UI_EXCLUDE_REASONS = {
    "login_fixture",
    "voice",
    "multi_finger",
    "scan_login",
    "factory_reset",
    "layout_review",
}

PURE_UI_EXCLUDE_SECOND_LEVEL = {
    "场景执行",
}

PURE_UI_EXCLUDE_FEATURE_KEYWORDS = {
    "熄火",
    "重启",
    "低电压",
    "高电压",
    "ACC",
    "D1",
    "场景执行",
    "触发与执行",
    "执行中",
}


PUBLIC_CONTROLS = {
    "Home1": {
        "control_name_cn": "HOME",
        "page": "公共控件",
        "location_desc": "中控底部 HOME 键",
        "image_path": "press_android/label_pic/T1V/Home1.png",
        "source_image": "press_android/label_pic/T1V/Home1.png",
        "notes": "系统 HOME 键",
    },
    "custom_scene_app_entry": {
        "control_name_cn": "自定义场景入口",
        "page": "公共控件",
        "location_desc": "应用入口 / Dock 图标",
        "image_path": "press_android/label_pic/custom_scene_app_entry.png",
        "source_image": "press_android/label_pic/custom_scene_app_entry.png",
        "notes": "进入自定义场景模块的入口图",
    },
}


ALIASES = {
    "自定义场景入口图标": "custom_scene_app_entry",
    "自定义场景应用图标": "custom_scene_app_entry",
    "底部dock栏将自定义场景图标": "custom_scene_app_entry",
    "自定义场景首页": "custom_scene_home_tabs",
    "自定义场景APP": "custom_scene_home_tabs",
    "推荐场景详情页": "recommend_detail_page",
    "使用此模版创建按钮": "recommend_detail_apply_btn",
    "使用此模板创建按钮": "recommend_detail_apply_btn",
    "试用按钮": "recommend_detail_try_btn",
    "返回图标": "recommend_detail_back",
    "场景推荐页签": "custom_scene_tab_recommend",
    "我的场景页签": "custom_scene_tab_mine",
    "新建场景页面": "new_scene_main_page",
    "新建场景入口": "new_scene_entry",
    "添加状态条件": "add_state_condition_btn",
    "添加动作": "add_action_btn",
    "条件列表页": "state_condition_dialog",
    "动作列表页": "action_dialog",
    "添加状态条件标题": "state_condition_title",
    "添加动作标题": "action_title",
    "全部条件": "all_condition_tab",
    "任意条件": "any_condition_tab",
    "保存": "save_btn",
    "取消": "cancel_btn",
}


def load_pipeline():
    spec = importlib.util.spec_from_file_location("custom_scene_cfg", CONFIG_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.PIPELINE


def load_cases():
    cases = {}
    with ANALYSIS_CSV.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            reasons = set(filter(None, row["reasons"].split(",")))
            if row["status"] not in {"draft", "ready"}:
                continue
            if reasons & PURE_UI_EXCLUDE_REASONS:
                continue
            if row["二级功能"] in PURE_UI_EXCLUDE_SECOND_LEVEL:
                continue
            feature_text = "\n".join([row["功能点"], row["前提条件"], row["操作步骤"], row["期望结果"]])
            if any(keyword in feature_text for keyword in PURE_UI_EXCLUDE_FEATURE_KEYWORDS):
                continue
            if row["用例序号"].startswith("T1V-S"):
                continue
            case_id = row["用例序号"]
            cases[case_id] = {
                "case_id": case_id,
                "row": row["row"],
                "second_level": row["二级功能"],
                "third_level": row["三级功能"],
                "feature": row["功能点"],
                "status": row["status"],
                "reasons": row["reasons"],
                "precondition": row["前提条件"],
                "steps": row["操作步骤"],
                "expected": row["期望结果"],
            }
    return cases


def load_detail_rows(valid_case_ids):
    rows = []
    with DETAIL_CSV.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            if row["case_id"] in valid_case_ids:
                rows.append(row)
    return rows


def build_control_index(pipeline):
    index = {}
    for item in pipeline["label_mappings"]:
        index[item["ascii_label"]] = {
            "control_name_cn": item["control_name_cn"],
            "page": item["page"],
            "location_desc": item.get("notes", ""),
            "image_path": f'{pipeline["label_root"]}/{item["ascii_label"]}.png',
            "source_image": item["source_image"],
            "notes": item.get("notes", ""),
        }

    existing_custom = ROOT / "press_android" / "label_pic" / "T1V" / "custom_scene"
    for img in existing_custom.glob("*.png"):
        index.setdefault(
            img.stem,
            {
                "control_name_cn": "",
                "page": "自定义场景",
                "location_desc": "",
                "image_path": f"press_android/label_pic/T1V/custom_scene/{img.name}",
                "source_image": "",
                "notes": "已有标签图，待补中文控件名",
            },
        )

    for label, meta in PUBLIC_CONTROLS.items():
        index[label] = meta
    return index


def build_usage_map(pipeline):
    usage_map = defaultdict(set)
    for case in pipeline["case_scripts"]:
        for line in case["steps"]:
            if line.startswith("touch: "):
                usage_map[line.split(": ", 1)[1].strip()].add("touch")
            elif line.startswith("check: "):
                usage_map[line.split(": ", 1)[1].strip()].add("check")
    return {k: "/".join(sorted(v)) for k, v in usage_map.items()}


def infer_label(normalized_name, index):
    if normalized_name in ALIASES:
        label = ALIASES[normalized_name]
        if label in index:
            return label, "alias"
    for label, meta in index.items():
        if meta["control_name_cn"] and meta["control_name_cn"] == normalized_name:
            return label, "exact_name"
    return "", "unmapped"


def build_breakdown_rows(cases, detail_rows, index, usage_map):
    rows = []
    case_stats = defaultdict(lambda: {"total": 0, "mapped": 0, "unmapped": 0})
    for item in detail_rows:
        label, match_type = infer_label(item["normalized_name"], index)
        meta = index.get(label, {})
        case_id = item["case_id"]
        case_stats[case_id]["total"] += 1
        if label:
            case_stats[case_id]["mapped"] += 1
        else:
            case_stats[case_id]["unmapped"] += 1
        image_path = meta.get("image_path", "")
        has_image = "yes" if (image_path and (ROOT / image_path).exists()) else "no"
        rows.append(
            {
                "case_id": case_id,
                "source_field": item["source_field"],
                "group": item["group"],
                "normalized_name": item["normalized_name"],
                "source_sentence": item["source_sentence"],
                "ascii_label": label,
                "usage_type": usage_map.get(label, ""),
                "mapping_status": match_type,
                "control_name_cn": meta.get("control_name_cn", ""),
                "page": meta.get("page", ""),
                "location_desc": meta.get("location_desc", ""),
                "image_path": image_path,
                "has_image": has_image,
                "source_image": meta.get("source_image", ""),
                "notes": meta.get("notes", ""),
            }
        )
    return rows, case_stats


def fit_image(image_path: Path, max_width=150, max_height=96):
    with PILImage.open(image_path) as img:
        width, height = img.size
    scale = min(max_width / width, max_height / height, 1.0)
    xl_img = XLImage(str(image_path))
    xl_img.width = int(width * scale)
    xl_img.height = int(height * scale)
    return xl_img, xl_img.height


def style_headers(ws, headers, widths):
    ws.freeze_panes = "A2"
    ws.append(headers)
    for idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]


def write_case_summary(ws, cases, case_stats):
    headers = ["用例ID", "二级功能", "三级功能", "功能点", "前提条件", "操作步骤", "期望结果", "UI项总数", "已映射", "待补图"]
    widths = [18, 18, 18, 28, 32, 38, 38, 10, 10, 10]
    style_headers(ws, headers, widths)
    row_idx = 2
    for case_id in sorted(cases.keys()):
        case = cases[case_id]
        stats = case_stats.get(case_id, {"total": 0, "mapped": 0, "unmapped": 0})
        values = [
            case_id,
            case["second_level"],
            case["third_level"],
            case["feature"],
            case["precondition"],
            case["steps"],
            case["expected"],
            stats["total"],
            stats["mapped"],
            stats["unmapped"],
        ]
        for col_idx, value in enumerate(values, 1):
            ws.cell(row=row_idx, column=col_idx, value=value).alignment = Alignment(vertical="top", wrap_text=True)
        row_idx += 1


def write_breakdown(ws, rows):
    headers = ["用例ID", "来源", "分组", "句子", "控件名", "标签名", "是否用于touch/check", "映射状态", "是否已有图", "页面", "控件位置", "图片路径", "来源图", "备注"]
    widths = [18, 10, 12, 42, 20, 24, 18, 12, 10, 16, 24, 48, 48, 24]
    style_headers(ws, headers, widths)
    row_idx = 2
    for row in rows:
        values = [
            row["case_id"],
            row["source_field"],
            row["group"],
            row["source_sentence"],
            row["normalized_name"],
            row["ascii_label"],
            row["usage_type"],
            row["mapping_status"],
            row["has_image"],
            row["page"],
            row["location_desc"],
            row["image_path"],
            row["source_image"],
            row["notes"],
        ]
        for col_idx, value in enumerate(values, 1):
            ws.cell(row=row_idx, column=col_idx, value=value).alignment = Alignment(vertical="top", wrap_text=True)
        row_idx += 1


def write_control_catalog(ws, rows):
    headers = ["标签名", "控件名", "页面", "控件位置", "图片路径", "控件图", "关联用例数", "关联用例"]
    widths = [24, 20, 16, 24, 48, 18, 10, 40]
    style_headers(ws, headers, widths)
    grouped = {}
    for row in rows:
        label = row["ascii_label"]
        if not label:
            continue
        entry = grouped.setdefault(
            label,
            {
                "control_name_cn": row["control_name_cn"] or row["normalized_name"],
                "page": row["page"],
                "location_desc": row["location_desc"],
                "image_path": row["image_path"],
                "cases": set(),
            },
        )
        entry["cases"].add(row["case_id"])

    row_idx = 2
    for label in sorted(grouped.keys()):
        item = grouped[label]
        ws.cell(row=row_idx, column=1, value=label)
        ws.cell(row=row_idx, column=2, value=item["control_name_cn"])
        ws.cell(row=row_idx, column=3, value=item["page"])
        ws.cell(row=row_idx, column=4, value=item["location_desc"])
        ws.cell(row=row_idx, column=5, value=item["image_path"])
        ws.cell(row=row_idx, column=7, value=len(item["cases"]))
        ws.cell(row=row_idx, column=8, value=",".join(sorted(item["cases"])))
        for col in [1, 2, 3, 4, 5, 7, 8]:
            ws.cell(row=row_idx, column=col).alignment = Alignment(vertical="center", wrap_text=True)
        img_path = ROOT / item["image_path"]
        if img_path.exists():
            xl_img, img_height = fit_image(img_path)
            ws.add_image(xl_img, f"F{row_idx}")
            ws.row_dimensions[row_idx].height = max(78, img_height * 0.8)
        row_idx += 1


def main():
    pipeline = load_pipeline()
    cases = load_cases()
    detail_rows = load_detail_rows(set(cases.keys()))
    index = build_control_index(pipeline)
    usage_map = build_usage_map(pipeline)
    breakdown_rows, case_stats = build_breakdown_rows(cases, detail_rows, index, usage_map)

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "pure_ui_cases"
    write_case_summary(ws_summary, cases, case_stats)

    ws_breakdown = wb.create_sheet("case_ui_breakdown")
    write_breakdown(ws_breakdown, breakdown_rows)

    ws_catalog = wb.create_sheet("controls_catalog")
    write_control_catalog(ws_catalog, breakdown_rows)

    out_path = OUT_DIR / "custom_scene_pure_ui_case_mapping_no_vehicle_v2.xlsx"
    wb.save(out_path)
    print(f"xlsx={out_path.relative_to(ROOT).as_posix()}")
    print(f"case_count={len(cases)}")
    print(f"breakdown_rows={len(breakdown_rows)}")


if __name__ == "__main__":
    main()
