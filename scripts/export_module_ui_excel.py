#!/usr/bin/env python3
import importlib.util
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage


ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "out" / "custom-scene-analysis"
CONFIG_PATH = ROOT / "config" / "custom_scene_ui_pipeline_config.py"


PUBLIC_CONTROLS = [
    {
        "module": "public",
        "control_name_cn": "HOME",
        "ascii_label": "Home1",
        "page": "公共控件",
        "location_desc": "中控底部 HOME 键",
        "image_path": "press_android/label_pic/T1V/Home1.png",
        "source_image": "press_android/label_pic/T1V/Home1.png",
        "notes": "系统 HOME 键",
    },
    {
        "module": "public",
        "control_name_cn": "自定义场景入口",
        "ascii_label": "custom_scene_app_entry",
        "page": "公共控件",
        "location_desc": "应用入口 / Dock 图标",
        "image_path": "press_android/label_pic/custom_scene_app_entry.png",
        "source_image": "press_android/label_pic/custom_scene_app_entry.png",
        "notes": "进入自定义场景模块的入口图",
    },
]


def load_pipeline():
    spec = importlib.util.spec_from_file_location("custom_scene_cfg", CONFIG_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.PIPELINE


def build_rows(pipeline):
    usage_map = {}
    for case in pipeline["case_scripts"]:
        for line in case["steps"]:
            if line.startswith("touch: "):
                label = line.split(": ", 1)[1].strip()
                usage_map.setdefault(label, set()).add("touch")
            elif line.startswith("check: "):
                label = line.split(": ", 1)[1].strip()
                usage_map.setdefault(label, set()).add("check")

    rows = []
    for item in pipeline["label_mappings"]:
        usage = "/".join(sorted(usage_map.get(item["ascii_label"], [])))
        rows.append(
            {
                "module": pipeline["module"],
                "control_name_cn": item["control_name_cn"],
                "ascii_label": item["ascii_label"],
                "usage_type": usage,
                "page": item["page"],
                "location_desc": item.get("notes", ""),
                "image_path": f'{pipeline["label_root"]}/{item["ascii_label"]}.png',
                "source_image": item["source_image"],
                "notes": item.get("notes", ""),
            }
        )
    for item in PUBLIC_CONTROLS:
        usage = "/".join(sorted(usage_map.get(item["ascii_label"], [])))
        row = dict(item)
        row["usage_type"] = usage
        rows.append(row)
    return rows


def fit_image(image_path: Path, max_width=180, max_height=110):
    with PILImage.open(image_path) as img:
        width, height = img.size
    scale = min(max_width / width, max_height / height, 1.0)
    xl_img = XLImage(str(image_path))
    xl_img.width = int(width * scale)
    xl_img.height = int(height * scale)
    return xl_img, xl_img.height


def style_sheet(ws, title):
    ws.freeze_panes = "A2"
    headers = ["模块", "控件名", "标签名", "是否用于touch/check", "页面", "控件位置", "控件图", "图片路径", "来源图"]
    ws.append(headers)
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = [12, 20, 24, 18, 18, 28, 18, 52, 52]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.title = title


def write_rows(ws, rows):
    row_idx = 2
    for row in rows:
        image_abspath = ROOT / row["image_path"]
        ws.cell(row=row_idx, column=1, value=row["module"])
        ws.cell(row=row_idx, column=2, value=row["control_name_cn"])
        ws.cell(row=row_idx, column=3, value=row["ascii_label"])
        ws.cell(row=row_idx, column=4, value=row["usage_type"])
        ws.cell(row=row_idx, column=5, value=row["page"])
        ws.cell(row=row_idx, column=6, value=row["location_desc"])
        ws.cell(row=row_idx, column=8, value=row["image_path"])
        ws.cell(row=row_idx, column=9, value=row["source_image"])
        for col in [1, 2, 3, 4, 5, 6, 8, 9]:
            ws.cell(row=row_idx, column=col).alignment = Alignment(vertical="center", wrap_text=True)

        if image_abspath.exists():
            xl_img, img_height = fit_image(image_abspath)
            ws.add_image(xl_img, f"G{row_idx}")
            ws.row_dimensions[row_idx].height = max(84, img_height * 0.8)
        else:
            ws.cell(row=row_idx, column=7, value="图片缺失")
            ws.row_dimensions[row_idx].height = 42
        row_idx += 1


def main():
    pipeline = load_pipeline()
    rows = build_rows(pipeline)
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    module_groups = {}
    for row in rows:
        module_groups.setdefault(row["module"], []).append(row)

    for module, items in module_groups.items():
        ws = wb.create_sheet(title=module[:31])
        style_sheet(ws, module[:31])
        write_rows(ws, items)

    out_path = OUT_DIR / "module_ui_mapping.xlsx"
    wb.save(out_path)
    print(f"xlsx={out_path.relative_to(ROOT).as_posix()}")
    print(f"sheets={','.join(module_groups.keys())}")
    print(f"rows={len(rows)}")


if __name__ == "__main__":
    main()
