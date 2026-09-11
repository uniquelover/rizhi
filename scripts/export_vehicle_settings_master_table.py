#!/usr/bin/env python3
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "out" / "vehicle-settings-analysis"
CONTROL_CSV = OUT_DIR / "vehicle_settings_control_mapping.csv"
PAGE_CSV = OUT_DIR / "vehicle_settings_page_tree.csv"


MASTER_HEADERS = [
    "模块",
    "一级页面",
    "二级页面",
    "页面名",
    "页面键",
    "控件名",
    "控件别名",
    "用途",
    "对应用例",
    "是否已抓图",
    "图片路径",
    "图片来源组",
    "是否需要CAN",
    "CAN说明",
    "是否纯UI",
    "是否已验证",
    "范围状态",
    "备注",
]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def build_alias(control_name: str, image_path: str) -> str:
    aliases: list[str] = []
    if control_name == "Home":
        aliases.extend(["Home1", "T1VHome", "T1VHome1"])
    elif control_name == "车辆设置icon":
        aliases.extend(["T1VDock车辆设置", "设置车辆未选中"])
    elif control_name == "充电结束":
        aliases.append("能量管理充电已结束")
    elif control_name == "后视镜调节":
        aliases.append("N50后视镜调节")
    elif control_name == "车门":
        aliases.extend(["车门锁", "N50车门锁"])

    if image_path:
        stem = Path(image_path).stem
        if stem != control_name and stem not in aliases:
            aliases.append(stem)
    return ", ".join(aliases)


def normalize_usage(usage_type: str) -> str:
    if not usage_type:
        return ""
    parts = []
    for item in usage_type.split("/"):
        if item == "ctouch_check":
            parts.append("check")
        elif item == "ctouch_touch":
            parts.append("touch")
        else:
            parts.append(item)
    deduped = []
    for item in parts:
        if item not in deduped:
            deduped.append(item)
    return "/".join(deduped)


def infer_can(control_name: str, used_by_cases: str) -> tuple[str, str]:
    can_keywords = ["充电", "锁车", "解锁", "闭锁", "后视镜", "位置灯", "近光灯", "后雾灯"]
    if any(keyword in control_name for keyword in can_keywords):
        return "是", "历史脚本或场景可能依赖 CAN 状态，需结合用例再确认。"
    if any(keyword in used_by_cases for keyword in ["车控车设", "能量管理"]):
        return "待确认", "当前来自车辆设置历史用例，是否依赖 CAN 需逐条确认。"
    return "否", ""


def infer_pure_ui(control_name: str, used_by_cases: str, can_flag: str) -> str:
    high_risk_keywords = ["恢复出厂设置"]
    if any(keyword in control_name for keyword in high_risk_keywords):
        return "否"
    if can_flag == "是":
        return "待确认"
    if used_by_cases:
        return "是"
    return "待确认"


def infer_verified(row: dict[str, str]) -> str:
    if row.get("scope_status") == "in_scope" and row.get("has_image") == "yes":
        return "已补图"
    return "待验证"


def build_master_rows() -> list[dict[str, str]]:
    page_rows = {row["page_key"]: row for row in read_csv(PAGE_CSV)}
    control_rows = read_csv(CONTROL_CSV)
    master_rows: list[dict[str, str]] = []

    for row in control_rows:
        page = page_rows.get(row.get("page_key", ""), {})
        can_flag, can_note = infer_can(row.get("control_name", ""), row.get("used_by_cases", ""))
        master_rows.append(
            {
                "模块": "车辆设置",
                "一级页面": row.get("level1", "") or page.get("level1", ""),
                "二级页面": row.get("level2", "") or page.get("level2", ""),
                "页面名": row.get("page_name", "") or page.get("page_name", ""),
                "页面键": row.get("page_key", ""),
                "控件名": row.get("control_name", ""),
                "控件别名": build_alias(row.get("control_name", ""), row.get("image_path", "")),
                "用途": normalize_usage(row.get("usage_type", "")),
                "对应用例": row.get("used_by_cases", ""),
                "是否已抓图": "是" if row.get("has_image") == "yes" else "否",
                "图片路径": row.get("image_path", ""),
                "图片来源组": row.get("image_group", ""),
                "是否需要CAN": can_flag,
                "CAN说明": can_note,
                "是否纯UI": infer_pure_ui(row.get("control_name", ""), row.get("used_by_cases", ""), can_flag),
                "是否已验证": infer_verified(row),
                "范围状态": row.get("scope_status", ""),
                "备注": row.get("notes", ""),
            }
        )
    return master_rows


def style_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.append(MASTER_HEADERS)
    widths = [10, 10, 12, 16, 16, 22, 24, 14, 38, 10, 56, 18, 12, 30, 10, 12, 18, 36]
    for idx, header in enumerate(MASTER_HEADERS, 1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]


def write_xlsx(rows: list[dict[str, str]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    style_sheet(ws)
    row_no = 2
    for row in rows:
        for col_no, field in enumerate(MASTER_HEADERS, 1):
            ws.cell(row=row_no, column=col_no, value=row.get(field, ""))
            ws.cell(row=row_no, column=col_no).alignment = Alignment(vertical="center", wrap_text=True)
        row_no += 1
    out_path = OUT_DIR / "车辆设置_总表.xlsx"
    wb.save(out_path)
    return out_path


def write_csv(rows: list[dict[str, str]]) -> Path:
    out_path = OUT_DIR / "车辆设置_总表.csv"
    with out_path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=MASTER_HEADERS)
        writer.writeheader()
        writer.writerows(rows)
    return out_path


def main() -> None:
    rows = build_master_rows()
    xlsx_path = write_xlsx(rows)
    csv_path = write_csv(rows)
    print(f"xlsx={xlsx_path.relative_to(ROOT).as_posix()}")
    print(f"csv={csv_path.relative_to(ROOT).as_posix()}")
    print(f"rows={len(rows)}")


if __name__ == "__main__":
    main()
